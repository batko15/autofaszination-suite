"""AutoFaszination Suite — Parser für Excel/JSON-Daten (Tuning-Listen & Markenhäuser).

Unterstützte Formate: .json (Liste von Objekten) und .xlsx (erstes Tabellenblatt,
Spaltenüberschriften in der ersten Zeile). Akzeptiert werden camelCase-, snake_case-
und deutsche Feldnamen.
"""
import io
import json
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

# ─── Feld-Aliase ─────────────────────────────────────────────────────────────

VEHICLE_FIELDS: Dict[str, List[str]] = {
    "brand": ["brand", "marke", "Brand"],
    "model": ["model", "modell", "Model"],
    "engine": ["engine", "motor", "motorisierung", "Engine"],
    "fuel": ["fuel", "kraftstoff", "Fuel"],
    "euro_norm": ["euroNorm", "euro_norm", "euronorm", "euro"],
    "years": ["years", "jahr", "jahre", "baujahr"],
    "hp_orig": ["hpOrig", "hp_orig", "psOrig", "ps_orig", "psOriginal", "ps"],
    "nm_orig": ["nmOrig", "nm_orig", "psOriginalNm", "nmOriginal", "nm"],
    "hp_tuned": ["hpTuned", "hp_tuned", "psGetunt", "ps_getunt", "psTuned"],
    "nm_tuned": ["nmTuned", "nm_tuned", "nmGetunt", "nm_getunt", "nmTuned"],
    "fuel_saving": ["fuelSaving", "fuel_saving", "verbrauchErsparnis", "ersparnis", "fuelSavingPct"],
    "price": ["price", "preis", "preisChf", "uvp", "uvpBrutto", "priceChf"],
    "product_code": ["productCode", "product_code", "produkt", "code"],
    "install_min": ["installMin", "install_min", "einbauMin", "einbauzeit"],
}

PARTNER_FIELDS: Dict[str, List[str]] = {
    "company": ["company", "firma", "markenhaus", "name", "companyName"],
    "brands": ["brands", "marken", "brand", "marke"],
    "primary_brand": ["primaryBrand", "primary_brand", "hauptmarke"],
    "contact_person": ["contactPerson", "contact_person", "ansprechpartner", "ansprechperson", "kontakt"],
    "street": ["street", "strasse", "adresse"],
    "zip": ["zip", "plz", "zipCode", "postleitzahl"],
    "city": ["city", "ort", "stadt", "localite"],
    "canton": ["canton", "kanton", "kt"],
    "country": ["country", "land"],
    "phone": ["phone", "telefon", "tel", "telefonnummer"],
    "email": ["email", "mail", "eMail", "emailadresse"],
    "website": ["website", "web", "internet", "url"],
    "priority": ["priority", "prioritaet", "priorität", "prio"],
    "capacity": ["capacity", "kapazitaet", "kapazität", "capacityVehiclesPerMonth", "kapazitaetMonat"],
}


def _pick(row: Dict[str, Any], aliases: List[str], default: Any = None) -> Any:
    for key in list(row.keys()):
        for alias in aliases:
            if key.lower().replace("_", "").replace(" ", "") == alias.lower().replace("_", "").replace(" ", ""):
                value = row[key]
                if value is not None and value != "":
                    return value
    return default


def _to_int(value: Any, default: int = 0) -> int:
    try:
        return int(float(str(value).replace(",", ".").strip()))
    except (TypeError, ValueError):
        return default


def _to_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(str(value).replace(",", ".").replace("CHF", "").strip())
    except (TypeError, ValueError):
        return default


def _to_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_upload(filename: str, content: bytes) -> Tuple[List[Dict[str, Any]], List[str]]:
    """Liest .json oder .xlsx und gibt normalisierte Zeilen + Fehlermeldungen zurück."""
    errors: List[str] = []
    name = (filename or "").lower()
    if name.endswith(".json"):
        try:
            data = json.loads(content.decode("utf-8-sig"))
            if isinstance(data, dict):
                data = data.get("vehicles") or data.get("partners") or data.get("fahrzeuge") or data.get("markenhaeuser") or [data]
            rows = data if isinstance(data, list) else []
        except Exception as exc:
            return [], [f"JSON konnte nicht gelesen werden: {exc}"]
    elif name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            raw_rows = list(ws.iter_rows(values_only=True))
            wb.close()
            if not raw_rows:
                return [], ["Die Excel-Datei enthält keine Zeilen."]
            headers = [_to_str(h) for h in raw_rows[0]]
            rows = []
            for r in raw_rows[1:]:
                if all(v is None or v == "" for v in r):
                    continue
                rows.append({headers[i]: r[i] for i in range(len(headers)) if i < len(r)})
        except Exception as exc:
            return [], [f"Excel konnte nicht gelesen werden: {exc}"]
    else:
        return [], ["Nicht unterstütztes Dateiformat — bitte .json oder .xlsx verwenden."]
    if not rows:
        errors.append("Keine Datensätze gefunden.")
    return rows, errors


# ─── Fahrzeuge normalisieren ─────────────────────────────────────────────────

def normalize_vehicle(row: Dict[str, Any]) -> Tuple[Optional[Dict[str, Any]], Optional[str]]:
    brand = _to_str(_pick(row, VEHICLE_FIELDS["brand"]))
    model = _to_str(_pick(row, VEHICLE_FIELDS["model"]))
    engine = _to_str(_pick(row, VEHICLE_FIELDS["engine"]))
    if not brand or not model:
        return None, "Marke und Modell sind erforderlich."
    hp_orig = _to_int(_pick(row, VEHICLE_FIELDS["hp_orig"]), -1)
    hp_tuned = _to_int(_pick(row, VEHICLE_FIELDS["hp_tuned"]), -1)
    if hp_orig <= 0 or hp_tuned <= 0:
        return None, "PS Original und PS Tuning sind erforderlich (Zahlen > 0)."
    if hp_tuned < hp_orig:
        return None, "Tuning-PS darf nicht kleiner als Original-PS sein."

    fuel_raw = _to_str(_pick(row, VEHICLE_FIELDS["fuel"], "benzin")).lower()
    fuel = "diesel" if "diesel" in fuel_raw else ("hybrid" if "hybrid" in fuel_raw else "benzin")

    product_raw = _to_str(_pick(row, VEHICLE_FIELDS["product_code"], "BA")).upper()
    product_code = product_raw if product_raw in ("DA", "BA", "K", "GA") else "BA"

    return {
        "brand": brand,
        "model": model,
        "engine": engine or f"{hp_orig} PS",
        "fuel": fuel,
        "euro_norm": _to_str(_pick(row, VEHICLE_FIELDS["euro_norm"])) or None,
        "years": _to_str(_pick(row, VEHICLE_FIELDS["years"])) or None,
        "hp_orig": hp_orig,
        "nm_orig": _to_int(_pick(row, VEHICLE_FIELDS["nm_orig"]), 0),
        "hp_tuned": hp_tuned,
        "nm_tuned": _to_int(_pick(row, VEHICLE_FIELDS["nm_tuned"]), 0),
        "fuel_saving": _to_int(_pick(row, VEHICLE_FIELDS["fuel_saving"], 8), 8),
        "price": _to_float(_pick(row, VEHICLE_FIELDS["price"], 1090.0), 1090.0),
        "product_code": product_code,
        "install_min": _to_int(_pick(row, VEHICLE_FIELDS["install_min"], 15), 15),
    }, None


# ─── Markenhäuser normalisieren ──────────────────────────────────────────────

def normalize_partner(row: Dict[str, Any]) -> Tuple[Optional[Dict[str, Any]], Optional[str]]:
    company = _to_str(_pick(row, PARTNER_FIELDS["company"]))
    zip_code = _to_str(_pick(row, PARTNER_FIELDS["zip"]))
    city = _to_str(_pick(row, PARTNER_FIELDS["city"]))
    if not company:
        return None, "Firmenname ist erforderlich."
    if not (zip_code.isdigit() and len(zip_code) == 4):
        return None, f"«{company}»: PLZ {zip_code!r} ist ungültig (4-stellig erwartet)."

    brands_raw = _pick(row, PARTNER_FIELDS["brands"], [])
    if isinstance(brands_raw, str):
        brands_raw = [b.strip() for b in brands_raw.replace(";", ",").split(",") if b.strip()]
    brands = [str(b) for b in brands_raw] if isinstance(brands_raw, list) else []

    priority_raw = _to_str(_pick(row, PARTNER_FIELDS["priority"], "B")).upper()
    priority = priority_raw if priority_raw in ("A", "B", "C") else "B"

    return {
        "company": company,
        "brands": brands,
        "primary_brand": _to_str(_pick(row, PARTNER_FIELDS["primary_brand"])) or (brands[0] if brands else None),
        "contact_person": _to_str(_pick(row, PARTNER_FIELDS["contact_person"])) or None,
        "street": _to_str(_pick(row, PARTNER_FIELDS["street"])) or None,
        "zip": zip_code,
        "city": city or "",
        "canton": _to_str(_pick(row, PARTNER_FIELDS["canton"])).upper() or None,
        "country": _to_str(_pick(row, PARTNER_FIELDS["country"], "CH")) or "CH",
        "phone": _to_str(_pick(row, PARTNER_FIELDS["phone"])) or None,
        "email": _to_str(_pick(row, PARTNER_FIELDS["email"])) or None,
        "website": _to_str(_pick(row, PARTNER_FIELDS["website"])) or None,
        "priority": priority,
        "capacity": _to_int(_pick(row, PARTNER_FIELDS["capacity"], 5), 5),
    }, None
